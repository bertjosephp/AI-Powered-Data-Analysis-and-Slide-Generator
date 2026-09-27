"""Renderer tests: reopen the generated .pptx and check structure, data and fit.

No GUI renderer is involved. Text fit is checked with a conservative width
estimate (average glyph ≈ 0.55 em), which is enough to catch a budget or box
size that can't hold its worst-case text.
"""

import io
import math
from datetime import date
from pathlib import Path

import pytest
from pptx import Presentation
from pptx.util import Emu, Pt

from app.schemas.deck import (
    ChartInsightSlide,
    ChartRef,
    ExecutiveSummarySlide,
    HypothesesSlide,
    HypothesisItem,
    Kpi,
    KpiSlide,
    MetricRef,
    NextStepsSlide,
    SlideSpec,
    Takeaway,
    TitleSlide,
)
from app.schemas.insights import Insights
from app.schemas.profile import DatasetProfile
from app.services.deck import fit
from app.services.deck.fit import fit_slides
from app.services.deck.pptx_renderer import PptxRenderer
from app.services.deck.resolve import resolve_slides
from app.services.deck.theme import DEFAULT_THEME
from app.services.eda.profiler import build_profile
from app.services.ingestion.loader import load_dataset

FIXTURES = Path(__file__).parent.parent / "fixtures"
AVG_GLYPH_EM = 0.55
LINE_HEIGHT = 1.2


@pytest.fixture(scope="module")
def profile() -> DatasetProfile:
    df = load_dataset((FIXTURES / "sample.csv").read_bytes(), "sample.csv", 10**8)
    return build_profile(df, 200_000)


def _render(specs: list[SlideSpec], profile: DatasetProfile) -> Presentation:
    slides = fit_slides(resolve_slides(specs, profile, "sample.csv", today=date(2026, 9, 26)))
    return Presentation(io.BytesIO(PptxRenderer().render(slides, "sample.csv")))


@pytest.fixture(scope="module")
def deck(profile: DatasetProfile) -> Presentation:
    insights = Insights.model_validate_json((FIXTURES / "llm_response.json").read_text())
    return _render(insights.slides, profile)


def _texts(slide) -> list[str]:  # type: ignore[no-untyped-def]
    return [s.text_frame.text for s in slide.shapes if s.has_text_frame and s.text_frame.text]


def _charts(slide):  # type: ignore[no-untyped-def]
    return [s.chart for s in slide.shapes if s.has_chart]


def test_deck_is_16_by_9_with_one_slide_per_spec(deck) -> None:  # type: ignore[no-untyped-def]
    assert (deck.slide_width, deck.slide_height) == (DEFAULT_THEME.width, DEFAULT_THEME.height)
    assert len(deck.slides) == 8


def test_title_slide_carries_resolved_meta(deck) -> None:  # type: ignore[no-untyped-def]
    texts = _texts(deck.slides[0])
    assert "What really drives our revenue" in texts
    assert "sample.csv · 40 rows × 9 columns · September 2026" in texts


def test_kpi_values_come_from_the_profile(deck, profile) -> None:  # type: ignore[no-untyped-def]
    texts = _texts(deck.slides[2])
    for expected in ["40", "463", "2,156", "34.2%", "Median order value", "median of revenue"]:
        assert expected in texts


def test_correlation_chart_is_native_with_profile_data(deck, profile) -> None:  # type: ignore[no-untyped-def]
    (chart,) = _charts(deck.slides[3])
    plot = chart.plots[0]
    assert list(plot.categories) == [f"{p.a} × {p.b}" for p in profile.top_correlations]
    assert list(plot.series[0].values) == [p.r for p in profile.top_correlations]
    assert not chart.has_legend


def test_top_values_chart_labels(deck) -> None:  # type: ignore[no-untyped-def]
    (chart,) = _charts(deck.slides[4])
    series = chart.plots[0].series[0]
    assert list(chart.plots[0].categories) == ["Gizmo", "Gadget", "Widget"]
    labels = [series.points[i].data_label.text_frame.text for i in range(3)]
    assert labels == ["15", "13", "12"]


def test_every_slide_has_speaker_notes(deck) -> None:  # type: ignore[no-untyped-def]
    assert all(s.notes_slide.notes_text_frame.text.strip() for s in deck.slides)


def test_footer_numbers_content_slides(deck) -> None:  # type: ignore[no-untyped-def]
    assert "3 / 8" in _texts(deck.slides[2])
    assert not any("/ 8" in t for t in _texts(deck.slides[0]))  # title slide has no footer


def test_unresolvable_chart_renders_text_only(profile) -> None:  # type: ignore[no-untyped-def]
    spec = ChartInsightSlide(
        layout="chart_insight",
        title="Nothing to chart",
        bullets=["One", "Two"],
        chart=ChartRef(chart="top_values", column="ghost", finding_id=None),
    )
    deck = _render([spec], profile)
    assert _charts(deck.slides[0]) == []
    assert {"One", "Two"} <= set(_texts(deck.slides[0]))


# ---------- geometry ----------


def _worst_case_deck(profile: DatasetProfile) -> Presentation:
    """Every text field at its full budget, with long words, at max item counts."""

    def text(n: int) -> str:
        words = "Mmmmmm wwwwww revenue discount returns product "
        return (words * 40)[:n]

    metric = MetricRef(metric="max", column="revenue", finding_id=None)
    specs: list[SlideSpec] = [
        TitleSlide(layout="title", title=text(fit.TITLE), subtitle=text(fit.SUBTITLE)),
        ExecutiveSummarySlide(
            layout="executive_summary",
            headline=text(fit.HEADLINE),
            takeaways=[Takeaway(title=text(fit.CARD_TITLE), text=text(fit.CARD_TEXT))] * 3,
        ),
        KpiSlide(
            layout="kpi_cards",
            title=text(fit.TITLE),
            kpis=[Kpi(label=text(fit.KPI_LABEL), metric=metric)] * 4,
        ),
        ChartInsightSlide(
            layout="chart_insight",
            title=text(fit.TITLE),
            bullets=[text(fit.BULLET)] * 3,
            chart=ChartRef(chart="correlations", column=None, finding_id=None),
        ),
        ChartInsightSlide(
            layout="chart_insight",
            title=text(fit.TITLE),
            bullets=[text(fit.BULLET)] * 3,
            chart=ChartRef(chart="top_values", column="ghost", finding_id=None),
        ),
        HypothesesSlide(
            layout="hypotheses",
            title=text(fit.TITLE),
            items=[
                HypothesisItem(
                    statement=text(fit.STATEMENT), test=text(fit.STATEMENT), confidence="high"
                )
            ]
            * 3,
        ),
        NextStepsSlide(layout="next_steps", title=text(fit.TITLE), steps=[text(fit.STEP)] * 4),
    ]
    return _render(specs, profile)


def _estimated_height(text: str, font_pt: float, box_width: Emu, spacing: float) -> float:
    chars_per_line = max(1, int(box_width / Pt(font_pt * AVG_GLYPH_EM)))
    lines = 0
    for paragraph in text.split("\n"):
        # word wrap: count lines by greedy fill
        line_len, n = 0, 1
        for word in paragraph.split(" "):
            add = len(word) + (1 if line_len else 0)
            if line_len + add > chars_per_line:
                n += 1 + (len(word) - 1) // chars_per_line
                line_len = len(word) % chars_per_line or chars_per_line
            else:
                line_len += add
        lines += n
    return lines * Pt(font_pt) * LINE_HEIGHT * max(spacing, 1.0)


def test_all_shapes_stay_on_the_slide(profile) -> None:  # type: ignore[no-untyped-def]
    deck = _worst_case_deck(profile)
    for i, slide in enumerate(deck.slides):
        for shape in slide.shapes:
            assert shape.left >= 0 and shape.top >= 0, (i, shape.name)
            assert shape.left + shape.width <= deck.slide_width, (i, shape.name)
            assert shape.top + shape.height <= deck.slide_height, (i, shape.name)


def test_worst_case_text_fits_its_box(profile) -> None:  # type: ignore[no-untyped-def]
    deck = _worst_case_deck(profile)
    overflows = []
    for i, slide in enumerate(deck.slides):
        for shape in slide.shapes:
            if not shape.has_text_frame or not shape.text_frame.text:
                continue
            paragraph = shape.text_frame.paragraphs[0]
            size = paragraph.runs[0].font.size.pt
            needed = _estimated_height(
                shape.text_frame.text, size, shape.width, paragraph.line_spacing or 1.0
            )
            if needed > shape.height * 1.02:
                overflows.append(
                    f"slide {i + 1}: {shape.text_frame.text[:30]!r} "
                    f"needs {needed / 914400:.2f}in, box {shape.height / 914400:.2f}in"
                )
    assert not overflows, "\n".join(overflows)


def test_estimator_sanity() -> None:
    one_line = _estimated_height("short", 12, Emu(914400 * 5), 1.0)
    assert math.isclose(one_line, Pt(12) * LINE_HEIGHT)
    assert _estimated_height("word " * 200, 12, Emu(914400 * 2), 1.0) > one_line * 10
