"""Slides that reference findings: resolution, native rendering, and fit."""

import io
from functools import cache
from pathlib import Path

from pptx import Presentation
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Pt

from app.schemas.deck import (
    ChartInsightSlide,
    ChartRef,
    Kpi,
    KpiSlide,
    MetricRef,
    ResolvedChartInsightSlide,
    ResolvedKpiSlide,
    SlideSpec,
)
from app.schemas.findings import Finding
from app.schemas.profile import DatasetProfile
from app.services.analysis.battery import run_battery
from app.services.deck.fit import fit_slides
from app.services.deck.pptx_renderer import PptxRenderer
from app.services.deck.resolve import finding_caption, resolve_slides
from app.services.eda.profiler import build_profile
from app.services.ingestion.loader import load_dataset

DATA = Path(__file__).parents[2] / "app" / "sample_data"


@cache
def _ecommerce() -> tuple[DatasetProfile, tuple[Finding, ...]]:
    df = load_dataset((DATA / "ecommerce_orders.csv").read_bytes(), "e.csv", 10**8)
    profile = build_profile(df, 200_000)
    return profile, tuple(run_battery(df, profile).findings)


def _find(kind: str) -> Finding:
    return next(f for f in _ecommerce()[1] if f.kind == kind)


def _chart_slide(finding_id: str) -> ChartInsightSlide:
    return ChartInsightSlide(
        layout="chart_insight",
        title="A finding",
        bullets=["Why it matters."],
        chart=ChartRef(chart="finding", column=None, finding_id=finding_id),
    )


def _render(specs: list[SlideSpec]) -> Presentation:
    profile, findings = _ecommerce()
    slides = fit_slides(resolve_slides(specs, profile, "e.csv", findings=list(findings)))
    return Presentation(io.BytesIO(PptxRenderer().render(slides, "e.csv")))


def test_finding_kpis_use_headlines_and_captions() -> None:
    profile, findings = _ecommerce()
    segment = _find("segment")
    spec = KpiSlide(
        layout="kpi_cards",
        title="Key numbers",
        kpis=[
            Kpi(
                label="Return gap",
                metric=MetricRef(metric="finding", column=None, finding_id=segment.id),
            ),
            Kpi(label="Ghost", metric=MetricRef(metric="finding", column=None, finding_id="F99")),
        ],
    )
    (resolved,) = resolve_slides([spec], profile, "e.csv", findings=list(findings))
    assert isinstance(resolved, ResolvedKpiSlide)
    assert [(k.value, k.caption) for k in resolved.kpis] == [
        (segment.headline, finding_caption(segment))  # the unknown F99 is dropped
    ]


def test_finding_charts_pick_a_style_by_kind() -> None:
    profile, findings = _ecommerce()
    specs: list[SlideSpec] = [_chart_slide(_find(k).id) for k in ("segment", "bins", "trend")]
    resolved = resolve_slides(specs, profile, "e.csv", findings=list(findings))
    styles = [
        s.chart.style for s in resolved if isinstance(s, ResolvedChartInsightSlide) and s.chart
    ]
    assert styles == ["bars", "columns", "line"]
    segment_chart = resolved[0].chart  # type: ignore[union-attr]
    assert segment_chart is not None
    assert segment_chart.values == _find("segment").chart.values
    assert segment_chart.reference == _find("segment").chart.reference


def test_native_chart_types_and_reference_note() -> None:
    deck = _render([_chart_slide(_find(k).id) for k in ("segment", "bins", "trend")])
    types = [next(s for s in slide.shapes if s.has_chart).chart.chart_type for slide in deck.slides]
    assert types == [
        XL_CHART_TYPE.BAR_CLUSTERED,
        XL_CHART_TYPE.COLUMN_CLUSTERED,
        XL_CHART_TYPE.LINE_MARKERS,
    ]
    texts = [s.text_frame.text for s in deck.slides[0].shapes if s.has_text_frame]
    assert any(t.startswith("Overall: ") for t in texts)
    line = next(s for s in deck.slides[2].shapes if s.has_chart).chart
    assert b"tickLblSkip" in line.category_axis._element.xml.encode()


def test_unknown_finding_chart_degrades_to_text() -> None:
    deck = _render([_chart_slide("F404")])
    assert not any(s.has_chart for s in deck.slides[0].shapes)


def test_long_kpi_headlines_shrink_to_fit() -> None:
    renderer = PptxRenderer()
    width = 2_000_000  # ~2.2in card interior
    assert renderer._kpi_size("40", 4, width) == 42
    long_size = renderer._kpi_size("17.6% vs 6.3%", 4, width)
    assert 18 <= long_size < 42
    # Estimated rendered width stays within the card.
    assert len("17.6% vs 6.3%") * 0.6 * Pt(long_size) <= width


def test_finding_bars_emphasize_the_standout() -> None:
    from app.services.deck.theme import DEFAULT_THEME

    segment = _find("segment")
    deck = _render([_chart_slide(segment.id)])
    chart = next(s for s in deck.slides[0].shapes if s.has_chart).chart
    colors = [str(p.format.fill.fore_color.rgb) for p in chart.plots[0].series[0].points]
    standout = max(range(len(segment.chart.values)), key=lambda i: abs(segment.chart.values[i]))
    assert colors[standout] == DEFAULT_THEME.accent.lstrip("#")
    assert all(
        c == DEFAULT_THEME.accent_muted.lstrip("#") for i, c in enumerate(colors) if i != standout
    )
